"""Excel 점검표. 시트: 요약, KISA 항목, 점검 결과(finding 전체), ISMS-P (초안), 증적 요청,
AI 해설 (참고, 해설이 있을 때만).

보안: 보고서에는 점검 대상 계정 안의 누군가가 정할 수 있는 값(ENI 설명, 태그, 리소스 이름 등)이
들어간다. openpyxl은 `=`로 시작하는 문자열을 수식으로 저장하므로, 모든 문자열 칸을 문자열
형식으로 강제해 수식 삽입(formula injection)을 막는다.
"""

import io
import json
from collections import Counter
from collections.abc import Sequence
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kompass.models import AuditReport, ItemOutcome
from kompass.reporters._format import (
    DISCLAIMER,
    EVIDENCE_NOTICE,
    EVIDENCE_UNVERIFIED,
    EXPLANATION_NOTICE,
    OUTCOME_LABELS,
    PRIORITY_LABELS,
    SCHEME_LABELS,
    STATUS_LABELS,
    STATUS_ORDER,
    evidence_summary,
    explanation_summary,
    investigation_summary,
    kst,
)

_EXCEL_CELL_LIMIT = 32767
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_OUTCOME_FILLS = {
    "취약": PatternFill("solid", fgColor="F8CBAD"),
    "확인 불가": PatternFill("solid", fgColor="F4B183"),
    "검토 필요": PatternFill("solid", fgColor="FFE699"),
    "양호": PatternFill("solid", fgColor="C6EFCE"),
}
_GRAY_FILL = PatternFill("solid", fgColor="E7E6E6")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _set(sheet: Any, row: int, column: int, value: Any) -> Any:
    cell = sheet.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = value[:_EXCEL_CELL_LIMIT]
        cell.data_type = "s"  # 수식으로 해석되지 않게 한다
    else:
        cell.value = value
    return cell


def _table(
    sheet: Any,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    widths: Sequence[int],
    outcome_column: str | None = None,
) -> None:
    for column, header in enumerate(headers, start=1):
        cell = _set(sheet, 1, column, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    outcome_index = headers.index(outcome_column) + 1 if outcome_column else None
    for r, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = _set(sheet, r, column, value)
            cell.alignment = _WRAP
            if column == outcome_index:
                cell.fill = _OUTCOME_FILLS.get(str(value), _GRAY_FILL)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"


def _summary(sheet: Any, report: AuditReport) -> None:
    scan, evaluation = report.scan, report.evaluation
    item_counts = Counter(i.outcome for i in report.kisa_items)
    finding_counts = Counter(f.status for f in report.findings)
    rows: list[tuple[str, Any]] = [
        ("주의", DISCLAIMER),
        ("계정", scan.account_id),
        ("파티션", scan.partition),
        ("점검 수행 주체", scan.caller_arn),
        ("점검 리전 (EC2·VPC)", ", ".join(scan.regions)),
        (
            "점검 제외 리전 (접근 거부, 범위 밖)",
            "; ".join(f"{r}: {m}" for r, m in scan.excluded_regions.items()) or "없음",
        ),
        ("점검 ID", scan.scan_id),
        ("수집 시각", f"{kst(scan.started_at)} ~ {kst(scan.finished_at)}"),
        ("판정 시각", kst(evaluation.evaluated_at)),
        ("보고서 생성", kst(report.generated_at)),
        ("kompass 버전", evaluation.kompass_version),
        ("매핑 데이터", report.mapping_digest),
        ("판정 기준값", json.dumps(evaluation.settings, ensure_ascii=False)),
        (
            "KISA 항목 (19개)",
            ", ".join(f"{OUTCOME_LABELS[o]} {item_counts.get(o, 0)}" for o in ItemOutcome),
        ),
        (
            "finding",
            ", ".join(f"{STATUS_LABELS[s]} {finding_counts.get(s, 0)}" for s in STATUS_ORDER),
        ),
        ("건너뛴 룰 (양호 아님)", ", ".join(evaluation.skipped_rules) or "없음"),
        ("증적 요청", evidence_summary(report)),
        ("AI 해설 (참고)", explanation_summary(report)),
        ("추가 조사 (참고)", investigation_summary(report)),
        *([("AI 해설 주의", EXPLANATION_NOTICE)] if report.explanations is not None else []),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        _set(sheet, r, 1, label).font = Font(bold=True)
        _set(sheet, r, 2, value).alignment = _WRAP
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 100


def render_xlsx(report: AuditReport) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = "요약"
    _summary(summary, report)

    _table(
        book.create_sheet("KISA 항목"),
        ["코드", "항목", "중요도", "분류", "결과", "룰", "finding", "ISMS-P (초안)", "원문 쪽"],
        [
            (
                i.code,
                i.title,
                i.importance,
                i.automation.value,
                OUTCOME_LABELS[i.outcome],
                ", ".join(i.rule_ids),
                ", ".join(f"{k} {v}" for k, v in sorted(i.status_counts.items())),
                ", ".join(i.isms_p),
                i.page,
            )
            for i in report.kisa_items
        ],
        [8, 36, 8, 9, 16, 30, 24, 22, 9],
        outcome_column="결과",
    )

    _table(
        book.create_sheet("점검 결과"),
        [
            "룰",
            "KISA",
            "ISMS-P (초안)",
            "결과",
            "심각도",
            "리소스",
            "리전",
            "내용",
            "증적",
            "수집 시각",
        ],
        [
            (
                f.rule_id,
                ", ".join(f.mappings.kisa) if f.mappings else "",
                ", ".join(f.mappings.isms_p) if f.mappings else "",
                STATUS_LABELS[f.status],
                f.severity.value,
                f.resource,
                f.region or "",
                f.message,
                json.dumps(f.evidence, ensure_ascii=False),
                kst(f.collected_at),
            )
            for f in report.findings
        ],
        [9, 8, 18, 11, 9, 40, 15, 50, 60, 22],
        outcome_column="결과",
    )

    _table(
        book.create_sheet("ISMS-P (초안)"),
        ["번호", "영역", "통제항목", "결과", "KISA 항목", "룰"],
        [
            (
                c.code,
                c.area,
                c.name,
                OUTCOME_LABELS[c.outcome],
                ", ".join(c.kisa_items),
                ", ".join(c.rule_ids),
            )
            for c in report.isms_p_controls
        ],
        [9, 26, 34, 18, 20, 40],
        outcome_column="결과",
    )

    if report.evidence_requests:
        _evidence_sheet(book.create_sheet("증적 요청"), report)

    if report.explanations is not None:
        _explanation_sheet(book.create_sheet("AI 해설 (참고)"), report)
        if report.explanations.assessment is not None:
            _assessment_sheet(book.create_sheet("AI 종합 평가 (참고)"), report)

    if report.investigations is not None:
        _investigation_sheet(book.create_sheet("추가 조사 (참고)"), report)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _explanation_sheet(sheet: Any, report: AuditReport) -> None:
    explained = report.explanations
    assert explained is not None
    rows: list[tuple[Any, ...]] = []
    for item in explained.explanations:
        draft = item.explanation
        rows.append(
            (
                item.rule_id,
                "검증됨" if item.verified else "검증되지 않음",
                PRIORITY_LABELS[draft.priority],
                len(item.finding_keys),
                draft.summary,
                draft.risk,
                draft.priority_reason,
                "\n".join(
                    f'{SCHEME_LABELS[c.scheme]} {c.code}: "{c.quote}"' for c in draft.citations
                ),
                "\n".join(f"{n}. {s}" for n, s in enumerate(draft.remediation.steps, start=1)),
                "\n".join(draft.remediation.aws_cli),
                draft.remediation.terraform or "",
                "\n".join(f"{n.finding_key}: {n.note}" for n in draft.resource_notes),
                "\n".join(item.problems),
                item.model + (" (대체)" if item.fallback_used else ""),
            )
        )
    rows += [
        (
            f.rule_id,
            "해설 실패",
            "",
            len(f.finding_keys),
            f.error,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        )
        for f in explained.failures
    ]
    _table(
        sheet,
        [
            "룰",
            "검증",
            "우선순위",
            "대상 finding",
            "요약",
            "위험",
            "우선순위 근거",
            "근거 인용",
            "조치 절차",
            "AWS CLI",
            "Terraform",
            "리소스별 메모",
            "검증 문제",
            "모델",
        ],
        rows,
        [9, 13, 9, 9, 40, 40, 30, 50, 50, 60, 50, 40, 40, 18],
    )


def _assessment_sheet(sheet: Any, report: AuditReport) -> None:
    explained = report.explanations
    assert explained is not None and explained.assessment is not None
    item = explained.assessment
    draft = item.assessment
    rows: list[tuple[Any, ...]] = [
        ("요약", 0, "검증됨" if item.verified else "검증되지 않음", "", draft.summary, "")
    ]
    rows += [
        (
            "위험 시나리오",
            n,
            scenario.title,
            PRIORITY_LABELS[scenario.impact],
            scenario.description,
            "\n".join(scenario.finding_keys),
        )
        for n, scenario in enumerate(draft.scenarios, start=1)
    ]
    rows += [
        ("조치 순서", n, step.rule_id, "", step.reason, "")
        for n, step in enumerate(draft.action_order, start=1)
    ]
    rows += [("검증 문제", n, "", "", p, "") for n, p in enumerate(item.problems, start=1)]
    _table(
        sheet,
        ["구분", "순번", "제목·룰", "영향", "내용", "관련 finding"],
        rows,
        [14, 6, 30, 8, 80, 60],
    )


def _investigation_sheet(sheet: Any, report: AuditReport) -> None:
    investigated = report.investigations
    assert investigated is not None
    rows: list[tuple[Any, ...]] = []
    for item in investigated.investigations:
        draft = item.investigation
        label = "근거 검증됨" if item.verified else "검증되지 않음"
        rows.append((item.rule_id, "요약", label, draft.summary, ""))
        rows += [
            (item.rule_id, "사실", "", fact.statement, ", ".join(fact.call_ids))
            for fact in draft.facts
        ]
        rows += [(item.rule_id, "확인할 질문", "", q, "") for q in draft.open_questions]
        rows += [
            (item.rule_id, "finding 메모", n.finding_key, n.note, "") for n in draft.finding_notes
        ]
        rows += [(item.rule_id, "검증 문제", "", p, "") for p in item.problems]
        rows += [
            (
                item.rule_id,
                "도구 호출",
                c.call_id,
                f"{c.tool} {json.dumps(c.input_raw, ensure_ascii=False)}",
                c.error or "정상",
            )
            for c in item.tool_calls
        ]
    rows += [(f.rule_id, "조사 실패", "", f.error, "") for f in investigated.failures]
    _table(sheet, ["룰", "구분", "표시", "내용", "근거·결과"], rows, [9, 14, 30, 90, 20])


# --- 증적 요청 (ADR-037) ---

_EVIDENCE_COLUMNS = ["번호", "KISA 항목", "항목명", "결과", "요청 증적", "확인 목적", "원문 근거"]
_EVIDENCE_WIDTHS = [9, 9, 22, 14, 50, 40, 40]


def _evidence_rows(report: AuditReport) -> list[tuple[Any, ...]]:
    return [
        (
            r.id,
            g.code,
            g.title,
            OUTCOME_LABELS[g.outcome],
            r.request,
            r.purpose,
            f'"{r.basis}" ({g.page}쪽)',
        )
        for g in report.evidence_requests
        for r in g.requests
    ]


def _target_rows(report: AuditReport) -> list[tuple[Any, ...]]:
    return [
        (g.code, t.rule_id, STATUS_LABELS[t.status], t.resource, t.region or "-", t.message)
        for g in report.evidence_requests
        for t in g.targets
    ]


_TARGET_COLUMNS = ["KISA 항목", "룰", "결과", "리소스", "리전", "내용"]
_TARGET_WIDTHS = [9, 9, 12, 40, 16, 70]


def _evidence_sheet(sheet: Any, report: AuditReport) -> None:
    reasons = {g.code: g.reason for g in report.evidence_requests}
    counts = Counter(g.code for g in report.evidence_requests for _ in g.targets)
    rows = [(*row, reasons[row[1]], counts.get(row[1], 0)) for row in _evidence_rows(report)]
    _table(
        sheet,
        [*_EVIDENCE_COLUMNS, "사유", "확인 대상 리소스"],
        rows,
        [*_EVIDENCE_WIDTHS, 40, 12],
        outcome_column="결과",
    )


def render_evidence_xlsx(report: AuditReport) -> bytes:
    """담당자에게 보내는 증적 요청서. 제출 현황을 적을 빈 칸을 둔다."""
    scan = report.scan
    book = Workbook()
    guide = book.active
    guide.title = "안내"
    notices = [EVIDENCE_NOTICE]
    if any(not g.verified for g in report.evidence_requests):
        notices.append(EVIDENCE_UNVERIFIED)
    rows: list[tuple[str, Any]] = [
        ("안내", " ".join(notices)),
        ("계정", scan.account_id),
        ("점검 ID", scan.scan_id),
        ("판정 시각", kst(report.evaluation.evaluated_at)),
        ("요청서 생성", kst(report.generated_at)),
        ("요청 범위", evidence_summary(report)),
        ("작성 방법", "요청 증적 시트의 담당자·제출일·제출 여부·비고 칸을 채워 회신"),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        _set(guide, r, 1, label).font = Font(bold=True)
        _set(guide, r, 2, value).alignment = _WRAP
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 100

    _table(
        book.create_sheet("요청 증적"),
        [*_EVIDENCE_COLUMNS, "담당자", "제출일", "제출 여부", "비고"],
        [(*row, "", "", "", "") for row in _evidence_rows(report)],
        [*_EVIDENCE_WIDTHS, 12, 12, 10, 30],
        outcome_column="결과",
    )
    _table(
        book.create_sheet("확인 대상"),
        _TARGET_COLUMNS,
        _target_rows(report),
        _TARGET_WIDTHS,
        outcome_column="결과",
    )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
