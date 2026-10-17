"""제출된 증적 문서: 요청 ID 폴더에서 파일을 읽고 텍스트를 로컬에서 추출한다 (ADR-049).

폴더 구조는 `<증적 폴더>/<요청 ID>/…`다 (예: `evidence/CA-02-1/권한 검토.pdf`). 요청 ID 폴더
아래는 하위 폴더까지 모두 읽는다. 카탈로그에 없는 폴더, 증적 폴더 바로 아래의 파일, 숨김 파일,
Office 잠금 파일(`~$…`), 심볼릭 링크는 읽지 않고 `ignored`로 알린다. 링크를 따라가지 않는 것은
증적 폴더 밖의 파일이 LLM으로 가지 않게 하기 위해서다.

원본 PDF·이미지는 LLM에 그대로 보내지 않는다. 마스킹을 거칠 수 있도록 텍스트만 뽑으며 (원칙 5),
텍스트를 뽑지 못한 파일은 이유(`problem`)를 남겨 사람이 검토하게 한다. 제출 파일은 신뢰할 수
없는 입력이므로 파서 오류는 모두 그 파일의 문제로 기록하고 전체를 멈추지 않는다.
"""

import hashlib
from collections.abc import Iterable, Iterator
from io import BytesIO
from pathlib import Path

from kompass.models import EvidenceFile, EvidenceFileKind, EvidenceSubmission

MAX_FILE_BYTES = 20 * 1024 * 1024
TEXT_SUFFIXES = frozenset({".txt", ".md", ".csv", ".json"})
TEXT_ENCODINGS = ("utf-8-sig", "cp949")  # 국내 Windows 엑셀이 내보낸 CSV는 cp949인 경우가 많다


class EvidenceFolderError(Exception):
    """증적 폴더를 읽을 수 없음 (사용자에게 보이는 메시지)."""


def load_submission(
    root: Path, request_ids: Iterable[str], *, max_bytes: int = MAX_FILE_BYTES
) -> EvidenceSubmission:
    """증적 폴더를 읽는다. 파일은 요청 ID, 경로 순이다."""
    if not root.is_dir():
        raise EvidenceFolderError(f"증적 폴더가 없습니다: {root}")
    known = {r.upper() for r in request_ids}
    files: list[EvidenceFile] = []
    ignored: list[str] = []
    for entry in sorted(root.iterdir()):
        relative = entry.name
        if entry.is_symlink() or not entry.is_dir() or entry.name.upper() not in known:
            ignored.append(relative)
            continue
        request_id = entry.name.upper()
        for path in _walk(entry, root, ignored):
            files.append(_read(path, root, request_id, max_bytes))
    files.sort(key=lambda f: (f.request_id, f.path))
    return EvidenceSubmission(files=tuple(files), ignored=tuple(sorted(ignored)))


def _walk(folder: Path, root: Path, ignored: list[str]) -> Iterator[Path]:
    for entry in sorted(folder.iterdir()):
        if entry.is_symlink() or entry.name.startswith((".", "~$")):
            ignored.append(entry.relative_to(root).as_posix())
        elif entry.is_dir():
            yield from _walk(entry, root, ignored)
        elif entry.is_file():
            yield entry


def _read(path: Path, root: Path, request_id: str, max_bytes: int) -> EvidenceFile:
    size = path.stat().st_size
    kind = _kind(path)
    text: str | None = None
    problem: str | None
    if size > max_bytes:
        digest = _sha256_file(path)
        problem = f"파일이 너무 큼 ({size:,}바이트, 상한 {max_bytes:,}바이트)"
    else:
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        if kind is None:
            problem = f"지원하지 않는 형식 ({path.suffix or '확장자 없음'}). 사람이 검토해야 함"
        else:
            text, problem = _extract(kind, data)
    return EvidenceFile(
        request_id=request_id,
        path=path.relative_to(root).as_posix(),
        sha256=digest,
        size=size,
        kind=kind,
        text=text,
        problem=problem,
    )


def _kind(path: Path) -> EvidenceFileKind | None:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return EvidenceFileKind.PDF
    if suffix == ".xlsx":
        return EvidenceFileKind.XLSX
    if suffix in TEXT_SUFFIXES:
        return EvidenceFileKind.TEXT
    return None


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _extract(kind: EvidenceFileKind, data: bytes) -> tuple[str | None, str | None]:
    """(텍스트, 문제). 둘 중 하나만 값이 있다."""
    if kind == EvidenceFileKind.TEXT:
        return _decode(data)
    try:
        if kind == EvidenceFileKind.XLSX:
            return _xlsx_text(data), None
        return _pdf_text(data)
    except Exception as exc:  # 신뢰할 수 없는 파일이라 파서 오류는 모두 그 파일의 문제다
        return None, f"파일을 읽지 못함 ({type(exc).__name__})"


def _decode(data: bytes) -> tuple[str | None, str | None]:
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding), None
        except UnicodeDecodeError:
            continue
    return None, "텍스트 인코딩을 판별하지 못함 (UTF-8·CP949가 아님)"


def _xlsx_text(data: bytes) -> str:
    """시트마다 `[시트: 이름]` 머리줄 뒤에 행을 탭으로 이어 붙인다. 빈 행·행 끝 빈 칸은 뺀다."""
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        sections = []
        for sheet in book.worksheets:
            lines = [f"[시트: {sheet.title}]"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    lines.append("\t".join(cells))
            sections.append("".join(f"{line}\n" for line in lines))
        return "\n".join(sections)
    finally:
        book.close()


def _pdf_text(data: bytes) -> tuple[str | None, str | None]:
    """쪽마다 `[쪽 N]` 머리줄을 붙인다. 텍스트가 없으면 스캔 문서로 본다."""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        return None, "암호가 걸린 PDF라 읽지 못함. 암호를 푼 파일을 받아야 함"
    pages = [(n, (page.extract_text() or "").strip()) for n, page in enumerate(reader.pages, 1)]
    if not any(text for _, text in pages):
        return None, "PDF에 텍스트가 없음 (스캔 문서로 보임). 사람이 검토해야 함"
    return "\n\n".join(f"[쪽 {n}]\n{text}" for n, text in pages), None
