"""Excel 점검표 리포터."""

import io
from typing import Any

from openpyxl import load_workbook

from kompass.models import Status
from kompass.reporters.xlsx_reporter import render_evidence_xlsx, render_xlsx
from tests.reporters.sample import DEFAULT, EXPLAINED, finding, sample_report

SHEETS = ["요약", "KISA 항목", "점검 결과", "ISMS-P (초안)", "증적 요청"]


def _book(report: Any = DEFAULT) -> Any:
    return load_workbook(io.BytesIO(render_xlsx(report)))


def _rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def test_sheets() -> None:
    assert _book().sheetnames == SHEETS


def test_summary_sheet() -> None:
    values = [v for row in _rows(_book()["요약"]) for v in row if v is not None]
    assert any("인증 심사 결과를 보장하지 않음" in str(v) for v in values)
    assert "123456789012" in values
    assert "2026-10-14 08:07:54 KST ~ 2026-10-14 08:07:54 KST" in values


def test_kisa_sheet_has_19_items_with_labels() -> None:
    sheet = _book()["KISA 항목"]
    rows = _rows(sheet)
    assert rows[0][:5] == ("코드", "항목", "중요도", "분류", "결과")
    assert len(rows) == 1 + 19
    ca09 = next(r for r in rows if r[0] == "CA-09")
    assert ca09[4] == "취약"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_findings_sheet_lists_every_finding() -> None:
    rows = _rows(_book()["점검 결과"])
    assert len(rows) == 1 + len(DEFAULT.findings)
    header = rows[0]
    sg = next(r for r in rows[1:] if r[header.index("리소스")] == "sg-1" and r[0] == "SG-001")
    assert sg[header.index("결과")] == "취약"
    assert sg[header.index("KISA")] == "CA-09"
    assert "open|ssh" in sg[header.index("증적")]


def test_isms_p_sheet_has_all_controls() -> None:
    rows = _rows(_book()["ISMS-P (초안)"])
    assert len(rows) == 1 + 101
    first = rows[1]
    assert first[0] == "1.1.1"
    assert first[1] == "관리체계 수립 및 운영"
    assert first[2] == "경영진의 참여"


def test_status_cells_are_colored() -> None:
    sheet = _book()["KISA 항목"]
    header = [c.value for c in sheet[1]]
    col = header.index("결과") + 1
    fills = {
        sheet.cell(row=r, column=col).value: sheet.cell(row=r, column=col).fill.fgColor.rgb
        for r in range(2, sheet.max_row + 1)
    }
    assert fills["취약"] != fills["양호"]


def test_formula_injection_is_neutralized() -> None:
    # 점검 대상 계정의 누군가가 정할 수 있는 값(ENI 설명, 태그 등)이 수식으로 실행되면 안 된다
    malicious = '=HYPERLINK("http://evil.example","click")'
    report = sample_report(
        finding("VPC-002", Status.REVIEW, "eni-1", "ap-northeast-2", malicious, {"d": malicious})
    )
    sheet = _book(report)["점검 결과"]
    cells = [c for row in sheet.iter_rows(min_row=2) for c in row if c.value == malicious]
    assert cells, "악성 문자열이 그대로 보존되어야 한다"
    assert all(c.data_type == "s" for c in cells)


def test_long_evidence_is_truncated_to_excel_limit() -> None:
    report = sample_report(finding("S3-004", Status.FAIL, "b", None, "m", {"x": "a" * 40000}))
    rows = _rows(_book(report)["점검 결과"])
    header = rows[0]
    assert len(rows[1][header.index("증적")]) <= 32767


# --- AI 해설 ---


def test_explanation_sheet_only_when_explained() -> None:
    assert "AI 해설 (참고)" not in _book().sheetnames
    assert _book(EXPLAINED).sheetnames == [*SHEETS, "AI 해설 (참고)", "AI 종합 평가 (참고)"]


def test_explanation_sheet_rows() -> None:
    rows = _rows(_book(EXPLAINED)["AI 해설 (참고)"])
    header, *body = rows
    assert header[:4] == ("룰", "검증", "우선순위", "대상 finding")
    assert [(r[0], r[1]) for r in body] == [
        ("SG-001", "검증됨"),
        ("S3-004", "검증되지 않음"),
        ("IAM-008", "해설 실패"),
    ]


def test_explanation_commands_are_not_formulas() -> None:
    sheet = _book(EXPLAINED)["AI 해설 (참고)"]
    cells = [
        c for row in sheet.iter_rows(min_row=2) for c in row if str(c.value).startswith("=cmd")
    ]
    assert cells and all(c.data_type == "s" for c in cells)


def test_summary_mentions_explanations() -> None:
    values = [v for row in _rows(_book(EXPLAINED)["요약"]) for v in row if v is not None]
    assert any("검증됨 1" in str(v) and "실패 1" in str(v) for v in values)


def test_assessment_sheet() -> None:
    book = _book(EXPLAINED)
    assert book.sheetnames[-1] == "AI 종합 평가 (참고)"
    rows = _rows(book["AI 종합 평가 (참고)"])
    assert rows[0][:2] == ("구분", "순번")
    kinds = [r[0] for r in rows[1:]]
    assert kinds == ["요약", "위험 시나리오", "조치 순서", "조치 순서", "조치 순서"]


def test_investigation_sheet() -> None:
    from tests.reporters.sample import INVESTIGATED

    book = _book(INVESTIGATED)
    assert book.sheetnames[-1] == "추가 조사 (참고)"
    kinds = [r[1] for r in _rows(book["추가 조사 (참고)"])[1:]]
    assert kinds == [
        "요약",
        "사실",
        "사실",
        "확인할 질문",
        "finding 메모",
        "도구 호출",
        "도구 호출",
        "조사 실패",
    ]


# --- 증적 요청 (ADR-037) ---


def test_evidence_sheet_lists_requests_with_reasons() -> None:
    rows = _rows(_book()["증적 요청"])
    assert rows[0][:4] == ("번호", "KISA 항목", "항목명", "결과")
    ca02 = [r for r in rows[1:] if r[1] == "CA-02"]
    assert [r[0] for r in ca02] == ["CA-02-1", "CA-02-2", "CA-02-3", "CA-02-4"]
    assert all(r[-1] == 1 for r in ca02)  # IAM-008 검토 필요 1건
    codes = {r[1] for r in rows[1:]}
    assert "CA-03" in codes  # full 항목이라도 룰을 건너뛰었으면 넣는다
    assert "CA-13" not in codes  # full 항목은 결과가 확정되면 빠진다


def test_evidence_workbook_has_blank_columns_and_targets() -> None:
    book = load_workbook(io.BytesIO(render_evidence_xlsx(DEFAULT)))
    assert book.sheetnames == ["안내", "요청 증적", "확인 대상"]
    guide = [v for row in _rows(book["안내"]) for v in row if v is not None]
    assert any("원문 대조 확인 전 초안" in str(v) for v in guide)
    requests = _rows(book["요청 증적"])
    assert requests[0][-4:] == ("담당자", "제출일", "제출 여부", "비고")
    assert all(r[-4:] == (None, None, None, None) for r in requests[1:])
    targets = _rows(book["확인 대상"])
    assert targets[1:] == [("CA-02", "IAM-008", "검토 필요", "account", "-", DEFAULT_IAM_008)]


def test_evidence_target_cells_are_not_formulas() -> None:
    report = sample_report(finding("IAM-007", Status.REVIEW, "=cmd|' /C calc'!A0"))
    book = load_workbook(io.BytesIO(render_evidence_xlsx(report)))
    [cell] = [c for row in book["확인 대상"].iter_rows(min_row=2) for c in row if c.column == 4]
    assert cell.data_type == "s" and cell.value.startswith("=")


DEFAULT_IAM_008 = "루트 비밀번호 사용 이력을 확인할 수 없다"
